"""TT46 independent oracle (python-docx, MIT): compares two .docx files paragraph by paragraph.
stdin: lines "original<TAB>candidate"; stdout: one JSON line per pair
{file, ok, paras, diffs:[...]} — compares style name, text, and per-character bold/italic/underline/size/font/color
(expanded per character so different run splits of the same formatting compare equal), tables cell text, section page size."""
import sys, json
import docx

def chars(p):
    out = []
    for r in p.runs:
        f = r.font
        key = (bool(r.bold), bool(r.italic), bool(r.underline), f.size.pt if f.size else None, f.name, str(f.color.rgb) if f.color is not None and f.color.type is not None and f.color.rgb is not None else None)
        out.extend((ch, key) for ch in r.text)
    return out

def summary(path):
    d = docx.Document(path)
    paras = [(p.style.name if p.style is not None else None, p.text, chars(p)) for p in d.paragraphs]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    secs = [(s.page_width, s.page_height) for s in d.sections]
    return paras, tables, secs

for line in sys.stdin:
    line = line.rstrip("\n")
    if not line: continue
    a, b = line.split("\t")
    try:
        pa, ta, sa = summary(a)
    except Exception as e:
        print(json.dumps({"file": a, "skip": f"python-docx cannot open original: {type(e).__name__}"})); continue
    try:
        pb, tb, sb = summary(b)
    except Exception as e:
        print(json.dumps({"file": a, "ok": False, "diffs": [f"candidate unreadable: {type(e).__name__}: {e}"]})); continue
    diffs = []
    if len(pa) != len(pb): diffs.append(f"paragraph count {len(pa)} != {len(pb)}")
    for i, (x, y) in enumerate(zip(pa, pb)):
        if x[0] != y[0]: diffs.append(f"p{i} style {x[0]!r} != {y[0]!r}")
        if x[1] != y[1]: diffs.append(f"p{i} text {x[1][:40]!r} != {y[1][:40]!r}")
        elif x[2] != y[2]:
            k = next(j for j in range(len(x[2])) if x[2][j] != y[2][j]) if len(x[2]) == len(y[2]) else -1
            diffs.append(f"p{i} formatting differs at char {k}: {x[2][k] if k >= 0 else len(x[2])} vs {y[2][k] if k >= 0 else len(y[2])}")
    if ta != tb: diffs.append("table cell text differs")
    if sa != sb: diffs.append("section page sizes differ")
    print(json.dumps({"file": a, "ok": not diffs, "paras": len(pa), "diffs": diffs[:5]}))
