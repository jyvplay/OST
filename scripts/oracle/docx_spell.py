#!/usr/bin/env python3
"""TT59 Writer proofing oracle (TEST-ONLY): builds the proofing fixture with python-docx + lxml (independent of Fusion's
writer) and reads saved files back with lxml.

    docx_spell.py make <out.docx>      fixture: misspellings, a tracked deletion, w:noProof, w:lang fr-FR, a table cell,
                                       a repeated word, URL / e-mail, ordinals, curly apostrophes
    docx_spell.py read <file.docx>     JSON {paras:[{text, ins, del}]} — text = what Word shows with changes pending
                                       (w:t outside w:del), ins = inserted text, del = deleted text (w:delText)
"""
import json, sys, copy
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

def make(out):
    d = Document()
    d.add_paragraph("Teh quick brown fox jumpd over the lazy dog.")
    p = d.add_paragraph("Visit https://exampel.com/pathh or mail someone@exampel.org about the 21st and 113th items, it\u2019s fine.")
    p = d.add_paragraph("")
    r = p.add_run("Fusionz "); rpr = r._r.get_or_add_rPr(); rpr.append(etree.SubElement(rpr, qn("w:noProof")))
    r = p.add_run("bonjour mademoiselle "); rpr = r._r.get_or_add_rPr(); lang = etree.SubElement(rpr, qn("w:lang")); lang.set(qn("w:val"), "fr-FR")
    p.add_run("and recieve this.")
    d.add_paragraph("This is is a repeated word and NASA plus XYZZY.")
    t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "Tabel cell"; t.cell(0, 1).text = "fine words"
    # a tracked deletion holding a misspelling: <w:del><w:r><w:delText>wrongg </w:delText></w:r></w:del>
    p = d.add_paragraph("Kept ")
    de = etree.SubElement(p._p, qn("w:del")); de.set(qn("w:id"), "901"); de.set(qn("w:author"), "Oracle"); de.set(qn("w:date"), "2026-10-22T00:00:00Z")
    rr = etree.SubElement(de, qn("w:r")); dt = etree.SubElement(rr, qn("w:delText")); dt.text = "wrongg "; dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r = etree.SubElement(p._p, qn("w:r")); tt = etree.SubElement(r, qn("w:t")); tt.text = "text."
    d.save(out)

def read(path):
    import zipfile
    x = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    out = []
    for p in x.iter(f"{{{W}}}p"):
        text, ins, dele = [], [], []
        for n in p.iter(f"{{{W}}}t", f"{{{W}}}delText"):
            anc = [a.tag for a in n.iterancestors()]
            if n.tag == f"{{{W}}}delText" or f"{{{W}}}del" in anc:
                dele.append(n.text or "")
                continue
            text.append(n.text or "")
            if f"{{{W}}}ins" in anc:
                ins.append(n.text or "")
        out.append({"text": "".join(text), "ins": "".join(ins), "del": "".join(dele)})
    print(json.dumps({"paras": out}, ensure_ascii=False))

if __name__ == "__main__":
    {"make": make, "read": read}[sys.argv[1]](sys.argv[2])
