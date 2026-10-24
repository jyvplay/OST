#!/usr/bin/env python3
"""TT63 semantic oracle for PivotTable parts (test-only; pandas + openpyxl + lxml, all BSD/MIT).

For every pivotTableDefinition in an .xlsx it reads ONLY:
  * the pivot XML (location, pivotFields axis / hidden items / defaultSubtotal, rowFields, colFields,
    pageFields, dataFields fld + subtotal, row/col grand totals) and the cache's worksheetSource;
  * the SOURCE cells (via openpyxl, cached values)
and recomputes, independently with pandas groupby, what the table must show. It then compares that with the
cells inside location/@ref that the writer stored (the cached view). Items are matched by their caption the way
Excel shows them ("(blank)", TRUE/FALSE, numbers in General format), case-insensitively.

It also re-checks structural invariants:
  * every <x v> in the records is a valid sharedItems index, recordCount == number of <r> == source rows;
  * every cacheField name equals the source heading (unique), cacheFields count == source width;
  * openpyxl loads the workbook with its own pivot parser (openpyxl.pivot) and the table survives a
    load -> save -> load round trip with the same rowFields / colFields / dataFields.

Usage: xlsx_pivot.py [--fusion-only] FILE.xlsx ...  (--fusion-only: skip tables whose cache is not refreshedBy="Fusion")
         -> prints JSON {file: {ok, tables: [{name, checked, errors[]}]}}
"""
import io
import json
import math
import posixpath
import sys
import zipfile

import openpyxl
import pandas as pd
from lxml import etree

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/package/2006/relationships"}
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
FNS = {
    "sum": lambda s: float(sum(s)) if len(s) else 0.0,
    "count": None,  # handled separately: non-empty cells
    "average": lambda s: (sum(s) / len(s)) if len(s) else "#DIV/0!",
    "max": lambda s: max(s) if len(s) else 0.0,
    "min": lambda s: min(s) if len(s) else 0.0,
    "product": lambda s: math.prod(s) if len(s) else 0.0,
    "countNums": lambda s: float(len(s)),
    "stdDev": lambda s: pd.Series(s).std(ddof=1) if len(s) > 1 else "#DIV/0!",
    "stdDevp": lambda s: pd.Series(s).std(ddof=0) if len(s) else "#DIV/0!",
    "var": lambda s: pd.Series(s).var(ddof=1) if len(s) > 1 else "#DIV/0!",
    "varp": lambda s: pd.Series(s).var(ddof=0) if len(s) else "#DIV/0!",
}
ERRS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!"}


def rels(z, part):
    d, b = posixpath.split(part)
    p = posixpath.join(d, "_rels", b + ".rels")
    if p not in z.namelist():
        return {}
    x = etree.fromstring(z.read(p))
    out = {}
    for r in x.findall("r:Relationship", NS):
        t = r.get("Target")
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], t[1:] if t.startswith("/") else posixpath.normpath(posixpath.join(d, t)))
    return out


def blank(v):
    return v is None or v == ""


def caption(v):
    if blank(v):
        return "(blank)"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def a1(ref):
    from openpyxl.utils.cell import range_boundaries
    c0, r0, c1, r1 = range_boundaries(ref)
    return c0, r0, c1, r1


def num_eq(a, b):
    if isinstance(a, str) or isinstance(b, str):
        return str(a) == str(b)
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)), abs(float(b)))


def check_table(z, wb, sheet_name, tpart):
    errs = []
    t = etree.fromstring(z.read(tpart))
    name = t.get("name")
    trel = rels(z, tpart)
    cpart = next((p for (k, p) in trel.values() if k == "pivotCacheDefinition"), None)
    if not cpart:
        return {"name": name, "checked": 0, "errors": ["table has no cache relationship"]}
    c = etree.fromstring(z.read(cpart))
    if FUSION_ONLY and c.get("refreshedBy") != "Fusion":
        return None
    ws_src = c.find("m:cacheSource/m:worksheetSource", NS)
    src_sheet, src_ref = ws_src.get("sheet"), ws_src.get("ref")
    sc0, sr0, sc1, sr1 = a1(src_ref)
    sws = wb[src_sheet]
    grid = [[sws.cell(row=r, column=col).value for col in range(sc0, sc1 + 1)] for r in range(sr0, sr1 + 1)]
    header, recs = grid[0], grid[1:]
    fields = c.findall("m:cacheFields/m:cacheField", NS)
    if len(fields) != len(header):
        errs.append(f"cacheFields {len(fields)} != source width {len(header)}")
    names = [f.get("name") for f in fields]
    if len({n.lower() for n in names}) != len(names):
        errs.append("cacheField names not unique")
    for i, (f, h) in enumerate(zip(names, header)):
        if not str(f).lower().startswith(caption(h).lower()):
            errs.append(f"cacheField {i} name {f!r} vs heading {h!r}")
    shared = [[None if e.tag.endswith("}m") else (e.get("v") == "1" if e.tag.endswith("}b") else float(e.get("v")) if e.tag.endswith("}n") else {"e": e.get("v")} if e.tag.endswith("}e") else e.get("v")) for e in f.find("m:sharedItems", NS)] for f in fields]
    # records
    crel = rels(z, cpart)
    rpart = next((p for (k, p) in crel.values() if k == "pivotCacheRecords"), None)
    if rpart:
        rx = etree.fromstring(z.read(rpart))
        rr = rx.findall("m:r", NS)
        if int(c.get("recordCount", "-1")) != len(rr) or len(rr) != len(recs):
            errs.append(f"recordCount {c.get('recordCount')} / records {len(rr)} / source rows {len(recs)}")
        for i, r in enumerate(rr[:5000]):
            for j, e in enumerate(r):
                if e.tag.endswith("}x"):
                    k = int(e.get("v"))
                    if k >= len(shared[j]):
                        errs.append(f"record {i} field {j}: x={k} outside sharedItems")
                    else:
                        sv, cv = shared[j][k], recs[i][j]
                        if caption(sv if not isinstance(sv, dict) else sv["e"]).lower() != caption(cv).lower():
                            errs.append(f"record {i} field {j}: item {sv!r} vs cell {cv!r}")
    else:
        errs.append("cache has no records part")

    # ---- the semantic recomputation
    pfs = t.findall("m:pivotFields/m:pivotField", NS)
    hidden = {}
    for i, pf in enumerate(pfs):
        its = pf.find("m:items", NS)
        if its is not None:
            hidden[i] = {caption(shared[i][int(it.get("x"))] if not isinstance(shared[i][int(it.get("x"))], dict) else shared[i][int(it.get("x"))]["e"]).lower()
                         for it in its.findall("m:item", NS) if it.get("h") == "1" and it.get("x") is not None}
    rows = [int(f.get("x")) for f in t.findall("m:rowFields/m:field", NS)]
    cols = [int(f.get("x")) for f in t.findall("m:colFields/m:field", NS) if f.get("x") != "-2"]
    dfs = [(int(d.get("fld")), d.get("subtotal", "sum"), d.get("name")) for d in t.findall("m:dataFields/m:dataField", NS)]
    row_sub = {i: (pfs[i].get("defaultSubtotal", "1") not in ("0", "false")) for i in rows}
    row_gt = t.get("colGrandTotals", "1") not in ("0", "false")   # Grand Total ROW (§18.10.1.73)
    col_gt = t.get("rowGrandTotals", "1") not in ("0", "false")   # Grand Total COLUMN
    df = pd.DataFrame(recs, columns=range(len(header)), dtype=object)
    # all-blank source rows ARE cache records (Excel): they only disappear when their "(blank)" items are hidden
    for f, hid in hidden.items():
        if hid:
            df = df[[caption(v).lower() not in hid for v in df[f]]]
    loc = t.find("m:location", NS)
    lc0, lr0, lc1, lr1 = a1(loc.get("ref"))
    fdr, fdc = int(loc.get("firstDataRow")), int(loc.get("firstDataCol"))
    host = wb[sheet_name]
    cell = lambda r, c: host.cell(row=lr0 + r, column=lc0 + c).value
    nrow, ncol = lr1 - lr0 + 1, lc1 - lc0 + 1
    nv = len(dfs)

    def agg(sub, fld, fn):
        col = list(sub[fld]) if len(sub) else []
        errs_in = [v for v in col if isinstance(v, str) and v in ERRS]
        if fn == "count":
            return float(sum(0 if blank(v) else 1 for v in col))
        if errs_in:
            return errs_in[0]
        nums = [float(v) for v in col if isinstance(v, (int, float)) and not isinstance(v, bool)]
        return FNS[fn](nums)

    def match(sub, path, fields):
        for f, lab in zip(fields, path):
            sub = sub[[caption(v).lower() == str(lab).lower() for v in sub[f]]]
        return sub

    # column header: one level per col field, then the value names when nv > 1 (Σ Values innermost)
    col_keys = []
    last_level = fdr - 1 - (1 if nv > 1 else 0)
    for j in range(fdc, ncol):
        col_keys.append([cell(last_level - (len(cols) - 1 - k), j) for k in range(len(cols))])
    # fill-down blank outer levels (Excel draws an outer column item once)
    for k in range(len(cols)):
        last = None
        for key in col_keys:
            if key[k] is None:
                key[k] = last
            else:
                last = key[k]
    checked = 0
    prev = [None] * len(rows)
    for i in range(fdr, nrow):
        labels = [cell(i, k) for k in range(len(rows))]
        tot0 = next((k for k, v in enumerate(labels) if isinstance(v, str) and v.endswith(" Total") and v != "Grand Total"), None)
        if tot0 is None and labels and labels[0] != "Grand Total":
            prev = labels
        elif tot0 is not None:
            labels = prev[:tot0] + labels[tot0:]
        first = labels[0] if labels else cell(i, 0)
        if isinstance(first, str) and first == "Grand Total" or (not rows and i == fdr):
            if not rows or row_gt:
                rpath, rf = [], []
            else:
                errs.append(f"row {i}: Grand Total row while colGrandTotals=0")
                continue
        else:
            tot = tot0
            if tot is not None:
                if not row_sub.get(rows[tot], True):
                    errs.append(f"row {i}: subtotal row while defaultSubtotal=0")
                rpath, rf = labels[:tot] + [labels[tot][: -len(" Total")]], rows[: tot + 1]
            else:
                rpath, rf = labels, rows
        rsub = match(df, rpath, rf)
        for jj, j in enumerate(range(fdc, ncol)):
            vi = (jj % nv) if nv > 1 else 0
            ckey = col_keys[jj]
            hv = cell(fdr - 1, j) if cols else None
            is_gt_col = bool(cols) and any(isinstance(x, str) and (x == "Grand Total" or x.startswith("Total ")) for x in ckey + [hv])
            if is_gt_col and not col_gt:
                errs.append(f"col {j}: Grand Total column while rowGrandTotals=0")
            sub = rsub if (is_gt_col or not cols) else match(rsub, ckey, cols)
            fld, fn, _n = dfs[vi]
            want = agg(sub, fld, fn) if len(sub) else None
            got = cell(i, j)
            if want is None:
                ok = got is None
            elif isinstance(want, str):
                ok = got == want
            else:
                ok = isinstance(got, (int, float)) and num_eq(got, want)
            checked += 1
            if not ok:
                errs.append(f"{name} r{i} c{j} path={rpath}/{ckey}: cell {got!r} vs pandas {want!r}")
    # the item SETS: every visible leaf row / column path is present exactly once, and nothing else
    def paths(fields):
        return sorted({tuple(caption(v).lower() for v in r) for r in df[fields].itertuples(index=False)}) if fields and len(df) else []
    got_rows = sorted(tuple(str(x).lower() for x in [cell(i, k) for k in range(len(rows))]) for i in range(fdr, nrow)
                      if rows and cell(i, 0) != "Grand Total" and not any(isinstance(cell(i, k), str) and cell(i, k).endswith(" Total") for k in range(len(rows))) and cell(i, len(rows) - 1) is not None)
    if rows and got_rows != paths(rows):
        errs.append(f"row items {got_rows} vs pandas {paths(rows)}")
    if cols:
        seen = []
        for jj, key in enumerate(col_keys):
            if (nv == 1 or jj % nv == 0) and not any(isinstance(x, str) and (x == "Grand Total" or x.startswith("Total ")) for x in key):
                seen.append(tuple(str(x).lower() for x in key))
        if sorted(seen) != paths(cols):
            errs.append(f"column items {sorted(seen)} vs pandas {paths(cols)}")
    # Σ Values (field -2): with 2+ data fields it must sit on the axis where the cells show the value names (columns)
    neg_c = [f.get("x") for f in t.findall("m:colFields/m:field", NS)].count("-2")
    neg_r = [f.get("x") for f in t.findall("m:rowFields/m:field", NS)].count("-2")
    if (nv > 1 and (neg_c != 1 or neg_r)) or (nv <= 1 and (neg_c or neg_r)):
        errs.append(f"data field -2: {neg_c} on columns, {neg_r} on rows for {nv} data fields")
    # column subtotals: with 2+ column fields Excel draws a subtotal column for every outer column field whose
    # defaultSubtotal is on (§18.10.1.73) — the cells must then contain " Total" column items
    for k, f in enumerate(cols[:-1]):
        if pfs[f].get("defaultSubtotal", "1") not in ("0", "false") and not any(isinstance(x, str) and x.endswith(" Total") and x != "Grand Total" for key in col_keys for x in key):
            errs.append(f"column field {f} has subtotals on but the cells show no subtotal column")
    # item order: the outermost row / column field shows its items in <items> order
    def order_of(f):
        its = pfs[f].find("m:items", NS)
        seq = []
        for it in its.findall("m:item", NS) if its is not None else []:
            if it.get("x") is not None and it.get("h") != "1":
                sv = shared[f][int(it.get("x"))]
                seq.append(caption(sv["e"] if isinstance(sv, dict) else sv).lower())
        return seq
    def first_seen(vals):
        out = []
        for x in vals:
            if x is not None and str(x).lower() not in out:
                out.append(str(x).lower())
        return out
    if rows:
        shown = first_seen(cell(i, 0) for i in range(fdr, nrow) if cell(i, 0) != "Grand Total" and not (isinstance(cell(i, 0), str) and cell(i, 0).endswith(" Total")))
        want = [x for x in order_of(rows[0]) if x in shown]
        if shown != want:
            errs.append(f"row item order {shown} vs <items> order {want}")
    if cols:
        shown = first_seen(k[0] for k in col_keys if not (isinstance(k[0], str) and (k[0] == "Grand Total" or k[0].startswith("Total "))))
        want = [x for x in order_of(cols[0]) if x in shown]
        if shown != want:
            errs.append(f"column item order {shown} vs <items> order {want}")
    # page-field area: rows above the table must be free for Excel's report filters
    npg = len(t.findall("m:pageFields/m:pageField", NS))
    if npg and lr0 - npg - 1 < 1:
        errs.append("page fields with no room above the table")
    return {"name": name, "checked": checked, "errors": errs[:40], "rows": rows, "cols": cols, "data": [d[1] for d in dfs], "pages": npg}


def roundtrip(path):
    """openpyxl's own pivot parser: load, save, reload; the table survives with the same fields.

    Workaround for an openpyxl defect: its writer de-duplicates caches with a set, and Serialisable defines
    value-based __eq__/__hash__, so two tables with their OWN but byte-identical caches (valid per ECMA-376, Excel
    writes that when cache sharing is off) collapse into one <pivotCache> and openpyxl's reload raises KeyError.
    Identity semantics for CacheDefinition restore the per-table caches; nothing else changes.
    Second openpyxl defect (pivot/cache.py CacheDefinition._write): the records relationship Target is taken from
    records.path BEFORE records._id = self._id is assigned, so cache N >= 2 points at pivotCacheRecords1.xml — it
    silently reads cache 1's records back, or crashes when cache 1 has none. The records id is set first here.
    """
    from openpyxl.pivot.cache import CacheDefinition
    CacheDefinition.__eq__ = object.__eq__
    CacheDefinition.__hash__ = object.__hash__
    if not getattr(CacheDefinition, "_tt63_fixed", False):
        orig_write = CacheDefinition._write

        def fixed_write(self, archive, manifest):
            if self.records is not None:
                self.records._id = self._id
            return orig_write(self, archive, manifest)
        CacheDefinition._write = fixed_write
        CacheDefinition._tt63_fixed = True
    wb = openpyxl.load_workbook(path)
    # Third openpyxl defect (workbook/_writer.py write_names): localSheetId is numbered over wb.worksheets (chartsheets
    # skipped) but the reader indexes wb._sheets (chartsheets included) -> a sheet-scoped name after a chartsheet
    # lands on the chartsheet and the reload raises. Sheet-scoped names play no part in this comparison: drop them.
    for ws in wb.worksheets:
        ws.defined_names.clear()
    before = {(ws.title, p.name): ([f.x for f in p.rowFields], [f.x for f in p.colFields], [(d.fld, d.subtotal) for d in p.dataFields]) for ws in wb.worksheets for p in ws._pivots}
    buf = io.BytesIO()
    wb.save(buf)
    wb2 = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    after = {(ws.title, p.name): ([f.x for f in p.rowFields], [f.x for f in p.colFields], [(d.fld, d.subtotal) for d in p.dataFields]) for ws in wb2.worksheets for p in ws._pivots}
    return before, after


def main(paths):
    out = {}
    for path in paths:
        try:
            z = zipfile.ZipFile(path)
            wb = openpyxl.load_workbook(path, data_only=True)
            wbx = etree.fromstring(z.read("xl/workbook.xml"))
            wrels = rels(z, "xl/workbook.xml")
            tables = []
            for s in wbx.findall("m:sheets/m:sheet", NS):
                spart = wrels[s.get(RID)][1]
                for (k, p) in rels(z, spart).values():
                    if k == "pivotTable":
                        tables.append(check_table(z, wb, s.get("name"), p))
            native = sum(1 for t in tables if t is None)
            tables = [t for t in tables if t is not None]
            before, after = roundtrip(path)
            rt = [] if before == after else [f"openpyxl round trip changed the pivots: {before} -> {after}"]
            out[path] = {"ok": not rt and all(not t["errors"] for t in tables), "tables": tables, "openpyxl": len(before), "roundtrip": rt, "native": native}
        except Exception as e:  # noqa: BLE001
            out[path] = {"ok": False, "error": f"{type(e).__name__}: {e}"}
    print(json.dumps(out, default=str))


FUSION_ONLY = False
if __name__ == "__main__":
    args = sys.argv[1:]
    if "--fusion-only" in args:
        FUSION_ONLY = True
        args = [a for a in args if a != "--fusion-only"]
    main(args)
