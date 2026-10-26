#!/usr/bin/env python3
"""TT75 oracle: Excel tables as openpyxl (MIT) reads them — independent of src/lib/formula.
Usage: xlsx_tables.py FILE... → JSON {file: {"tables": [...], "formulas": {sheet!A1: text}, "headers": {...}} | {"error": ...}}"""
import json, sys, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

def one(path):
    wb = load_workbook(path, data_only=False)
    tables, formulas, headers = [], {}, {}
    for ws in wb.worksheets:
        for t in ws.tables.values():
            cols = [c.name for c in t.tableColumns]
            af = t.autoFilter
            fcols = {}
            if af is not None:
                for fc in af.filterColumn or []:
                    vals = [f for f in (fc.filters.filter if fc.filters is not None else [])]
                    if fc.filters is not None and fc.filters.blank: vals.append("")
                    fcols[str(fc.colId)] = vals
            tables.append({"sheet": ws.title, "name": t.displayName or t.name, "ref": t.ref, "header": 0 if t.headerRowCount == 0 else 1,
                           "totals": int(t.totalsRowCount or 0), "columns": cols, "autoFilter": af.ref if af is not None else None, "filter": fcols, "id": t.id})
            # header cells: value + python type name
            from openpyxl.utils.cell import range_boundaries
            c0, r0, c1, r1 = range_boundaries(t.ref)
            if t.headerRowCount != 0:
                for c in range(c0, c1 + 1):
                    v = ws.cell(row=r0, column=c).value
                    headers[f"{ws.title}!{ws.cell(row=r0, column=c).coordinate}"] = [type(v).__name__, v if isinstance(v, (str, int, float, bool)) or v is None else str(v)]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "[" in c.value:
                    formulas[f"{ws.title}!{c.coordinate}"] = c.value
    return {"tables": tables, "formulas": formulas, "headers": headers}

out = {}
for p in sys.argv[1:]:
    try: out[p] = one(p)
    except Exception as e: out[p] = {"error": f"{type(e).__name__}: {e}"[:300]}
print(json.dumps(out, ensure_ascii=False))
