"""TT74 K1 oracle (test-only): python-pptx (MIT) reading of every deck: per slide (sldIdLst order) the hidden flag, the
paragraph texts of every text frame in shape-tree order (groups recursed; python-pptx does not read
mc:AlternateContent), table cells (spanned cells skipped), and the notes text.
usage: pptx_outline.py DIR  -> JSON {name: {"slides": [...]} | {"error": "..."}}"""
import json, os, sys
from pptx import Presentation


def paras(tf):
    return [p.text.replace("\v", "\n").replace("\x0b", "\n") for p in tf.paragraphs]


def walk(shapes, out, depth=0):
    if depth > 64:
        return
    for sh in shapes:
        if sh.shape_type is not None and sh.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            walk(sh.shapes, out, depth + 1)
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    if c.is_spanned:
                        continue
                    out.append("|" + c.text.replace("\v", "\n"))
            continue
        # has_text_frame is True for every p:sp and text_frame CREATES a txBody when absent: read only real ones
        if sh.has_text_frame and sh._element.tag.endswith("}sp") and sh._element.txBody is not None:
            out.extend(paras(sh.text_frame))


MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
SHAPES = {"sp", "grpSp", "graphicFrame", "cxnSp", "pic", "contentPart"}


def resolve_alt(src, dst):
    """copy of the deck with each mc:AlternateContent in a slide shape tree replaced by the shape of the branch
    PowerPoint 2007 reads (mc:Fallback, else mc:Choice) - the reading Fusion applies (pml.altContent)"""
    import zipfile
    from lxml import etree
    zin = zipfile.ZipFile(src)
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        data = zin.read(info.filename)
        if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(".xml"):
            root = etree.fromstring(data)
            for ac in list(root.iter(f"{{{MC}}}AlternateContent")):
                par = ac.getparent()
                if par is None or etree.QName(par).localname not in ("spTree", "grpSp"):
                    continue
                br = ac.find(f"{{{MC}}}Fallback")
                if br is None:
                    br = ac.find(f"{{{MC}}}Choice")
                shape = next((c for c in br if isinstance(c.tag, str) and etree.QName(c).localname in SHAPES), None) if br is not None else None
                if shape is not None:
                    par.replace(ac, shape)
                else:
                    par.remove(ac)
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(info, data)
    zout.close()


def main(d):
    rep = {}
    for name in sorted(os.listdir(d)):
        if not name.endswith(".pptx"):
            continue
        try:
            path = os.path.join(d, name)
            if os.environ.get("RESOLVE_ALT"):
                import tempfile
                tmp = os.path.join(tempfile.mkdtemp(), name)
                resolve_alt(path, tmp)
                path = tmp
            prs = Presentation(path)
            slides = []
            for s in prs.slides:
                out = []
                walk(s.shapes, out)
                notes = ""
                if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
                    notes = "\n".join(paras(s.notes_slide.notes_text_frame))
                slides.append({"hidden": s._element.get("show") in ("0", "false"), "paras": out, "notes": notes})
            rep[name] = {"slides": slides}
        except Exception as e:
            rep[name] = {"error": f"{type(e).__name__}: {str(e)[:160]}"}
    print(json.dumps(rep, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])
