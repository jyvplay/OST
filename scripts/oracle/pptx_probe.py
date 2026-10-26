#!/usr/bin/env python3
"""TT51 oracle: python-pptx 1.0.2 (MIT) reading of .pptx files -> JSON, independent of our TypeScript reader.
usage: pptx_probe.py file.pptx [...]   prints one JSON object per file (one line each)."""
import json, sys
from pptx import Presentation
from pptx.util import Emu

def geo(sh):
    try:
        return [int(sh.left) if sh.left is not None else None, int(sh.top) if sh.top is not None else None,
                int(sh.width) if sh.width is not None else None, int(sh.height) if sh.height is not None else None]
    except Exception:
        return None

RUNS = False

def font_color(f):
    try:
        return str(f.color.rgb) if f.color and f.color.type is not None and f.color.rgb is not None else None
    except Exception:
        return None

def paras(tf):
    if RUNS:
        return [[p.text, p.alignment and str(p.alignment).split(" ")[0], p.level,
                 [[r.text, r.font.bold, r.font.italic, r.font.underline, r.font.size and r.font.size.pt, font_color(r.font)] for r in p.runs]]
                for p in tf.paragraphs]
    return [p.text for p in tf.paragraphs]

def shape(sh):
    o = {"id": sh.shape_id, "name": sh.name, "geo": geo(sh), "type": str(sh.shape_type) if sh.shape_type is not None else None}
    if sh.is_placeholder:
        try:
            o["ph"] = [str(sh.placeholder_format.type).split(" ")[0], sh.placeholder_format.idx]
        except Exception:
            pass
    if sh.has_text_frame:
        o["text"] = paras(sh.text_frame)
    if getattr(sh, "has_table", False) and sh.has_table:
        o["table"] = [[paras(c.text_frame) for c in r.cells] for r in sh.table.rows]
    if sh.shape_type is not None and "GROUP" in str(sh.shape_type):
        o["kids"] = [shape(k) for k in sh.shapes]
    return o

def probe(path):
    try:
        p = Presentation(path)
    except Exception as e:
        return {"file": path, "error": type(e).__name__}
    out = {"file": path, "size": [int(p.slide_width or 0), int(p.slide_height or 0)], "slides": []}
    for s in p.slides:
        hidden = s._element.get("show") in ("0", "false")
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame is not None else None
        out["slides"].append({"partname": str(s.part.partname), "hidden": hidden, "layout": s.slide_layout.name, "notes": notes,
                              "shapes": [shape(sh) for sh in s.shapes]})
    out["layouts"] = sum(len(m.slide_layouts) for m in p.slide_masters)
    return out

args = sys.argv[1:]
if args and args[0] == "--runs":
    RUNS = True
    args = args[1:]
for f in args:
    print(json.dumps(probe(f), default=str))
