#!/usr/bin/env python3
"""TT49 independent conditional-format EVALUATION oracle (test-only, never shipped).

Independent of the app's TypeScript code path end to end:
  * rules are parsed by openpyxl (MIT) from the main schema (x14 extension rules are invisible to it);
  * cell values are the values cached in the file by its authoring application (openpyxl data_only=True);
  * relative references in formula rules are shifted by openpyxl's Translator, and expression / cellIs-formula
    rules are evaluated by the `formulas` package (a separate Python Excel-formula engine);
  * rule semantics are written from the specs (ECMA-376 Part 1 §18.3.1.10, MS-XLSX CT_Cfvo, Microsoft's
    "Working with conditional formatting" data-bar formula) in a different language.
stdin: .xlsx paths   stdout: JSON lines {"file", "rules": [{sheet, idx, priority, type, cells: {A1: value}} | {..., skip}]}
value: true (format rule matched), icon index (iconSet, after reverse), bar length 0..1 (dataBar), "#rrggbb" (colorScale).
idx is "<conditionalFormatting index>.<cfRule index>" in document order (the app's provenance key).
"""
import datetime, json, math, re, statistics, sys, warnings, zipfile, posixpath
warnings.filterwarnings("ignore")
import openpyxl, numpy as np, schedula as sh, formulas
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from lxml import etree

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
ICON_N = {"3Arrows": 3, "3ArrowsGray": 3, "3Flags": 3, "3TrafficLights1": 3, "3TrafficLights2": 3, "3Signs": 3, "3Symbols": 3, "3Symbols2": 3,
          "4Arrows": 4, "4ArrowsGray": 4, "4RedToBlack": 4, "4Rating": 4, "4TrafficLights": 4, "5Arrows": 5, "5ArrowsGray": 5, "5Rating": 5, "5Quarters": 5}
DEFAULT_PCT = {3: [0, 33, 67], 4: [0, 25, 50, 75], 5: [0, 20, 40, 60, 80]}
EPOCH = datetime.datetime(1899, 12, 30)


class Err(str):
    """An error value (distinct from text)."""


def doc_order(path):
    """sheet title -> {priority: "i.j"} from the raw XML (document order of top-level <conditionalFormatting>)."""
    z = zipfile.ZipFile(path)
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall(f"{{{PR}}}Relationship")}
    out = {}
    for s in wbx.iter(f"{{{S}}}sheet"):
        t = target.get(s.get(f"{{{R}}}id"), "")
        part = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join("xl", t))
        if part not in z.namelist():
            continue
        root = etree.fromstring(z.read(part))
        m = {}
        for i, cf in enumerate(root.findall(f"{{{S}}}conditionalFormatting")):
            for j, rule in enumerate(cf.findall(f"{{{S}}}cfRule")):
                m[int(rule.get("priority", "0"))] = (f"{i}.{j}", cf.get("sqref", ""))
        out[s.get("name")] = m
    return out


def excel_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        if float(v).is_integer() and abs(v) < 1e15:
            return str(int(v))
        return format(v, ".15g")
    return str(v)


def rank(v):
    return 0 if isinstance(v, (int, float)) and not isinstance(v, bool) else 1 if isinstance(v, str) else 2


def cmp(op, v, w):
    """Excel comparison; v None = blank (takes w's type)."""
    if isinstance(v, Err) or isinstance(w, Err):
        return False
    if v is None and w is None:
        v = w = 0  # blank = blank
    if v is None:
        v = "" if isinstance(w, str) else False if isinstance(w, bool) else 0
    if w is None:
        w = "" if isinstance(v, str) else False if isinstance(v, bool) else 0
    if rank(v) != rank(w):
        d = rank(v) - rank(w)
    elif isinstance(v, str):
        a, b = v.lower(), w.lower(); d = (a > b) - (a < b)
    else:
        d = (float(v) > float(w)) - (float(v) < float(w))
    return {"=": d == 0, "<>": d != 0, "<": d < 0, ">": d > 0, "<=": d <= 0, ">=": d >= 0}[op]


def to_val(c):
    v = c.value
    if c.data_type == "e":
        return Err(str(v))
    if isinstance(v, datetime.datetime):
        return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, datetime.date):
        return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, str) and v == "":
        return ""
    return v


def percentile_inc(s, p):
    r = max(0, min(100, p)) / 100 * (len(s) - 1); i = int(math.floor(r))
    return s[i] + (s[i + 1] - s[i]) * (r - i) if i + 1 < len(s) else s[i]


class Ctx:
    def __init__(self, wb):
        self.wb = wb; self.cache = {}

    def value(self, sheet, ref):
        ws = self.wb[sheet]
        return to_val(ws[ref])

    def evaluate(self, sheet, formula, origin, cell):
        """formula (no '=') relative to origin, evaluated at cell via Translator + formulas."""
        f = Translator("=" + formula, origin=origin).translate_formula(cell)
        col, row = openpyxl.utils.cell.coordinate_from_string(cell)
        # ROW(ref)/COLUMN(ref) need the reference itself, which `formulas` receives only as a value → literal
        f = re.sub(r"\bROW\(\s*\$?[A-Z]{1,3}\$?(\d+)\s*\)", lambda m: m.group(1), f, flags=re.I)
        f = re.sub(r"\bCOLUMN\(\s*\$?([A-Z]{1,3})\$?\d+\s*\)", lambda m: str(openpyxl.utils.column_index_from_string(m.group(1).upper())), f, flags=re.I)
        f = re.sub(r"\bROW\(\)", str(row), f, flags=re.I)
        f = re.sub(r"\bCOLUMN\(\)", str(openpyxl.utils.column_index_from_string(col)), f, flags=re.I)
        f = re.sub(r"\bTODAY\(\)", str((datetime.datetime.combine(datetime.date.today(), datetime.time()) - EPOCH).days), f, flags=re.I)
        fn = self.cache.get(f)
        if fn is None:
            fn = formulas.Parser().ast(f)[1].compile(); self.cache[f] = fn
        args = []
        for name in fn.inputs:
            sh_name, ref = (name.rsplit("!", 1) if "!" in name else (sheet, name))
            sh_name = sh_name.strip("'")
            ws = next((w for w in self.wb.worksheets if w.title.upper() == sh_name.upper()), None)
            if ws is None:
                raise ValueError("sheet " + sh_name)
            if ":" in ref:
                c0, r0, c1, r1 = range_boundaries(ref)
                if (r1 - r0 + 1) * (c1 - c0 + 1) > 200000:
                    r1 = min(r1, ws.max_row); c1 = min(c1, ws.max_column)
                arr = np.empty((r1 - r0 + 1, c1 - c0 + 1), dtype=object)
                for r in range(r0, r1 + 1):
                    for c in range(c0, c1 + 1):
                        arr[r - r0, c - c0] = self.fval(ws.cell(r, c))
                args.append(arr)
            else:
                args.append(self.fval(ws[ref]))
        out = fn(*args)
        if isinstance(out, np.ndarray):
            out = out.ravel()[0]
        if out is sh.EMPTY:
            return None
        from formulas.tokens.operand import XlError
        if isinstance(out, XlError):
            return Err(str(out))
        if isinstance(out, np.bool_):
            return bool(out)
        if isinstance(out, np.number):
            return float(out)
        return out

    @staticmethod
    def fval(c):
        v = to_val(c)
        if v is None:
            return sh.EMPTY
        if isinstance(v, Err):
            from formulas.tokens.operand import XlError
            return XlError(str(v))
        return v


def truthy(v):
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (int, float, np.number)) and not isinstance(v, bool):
        return float(v) != 0
    return False


def const(f):
    t = f.strip()
    if len(t) >= 2 and t[0] == '"' and t[-1] == '"':
        return t[1:-1].replace('""', '"')
    try:
        return float(t)
    except ValueError:
        pass
    if t.upper() in ("TRUE", "FALSE"):
        return t.upper() == "TRUE"
    return None


def period_match(p, v):
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return False
    d = (EPOCH + datetime.timedelta(days=math.floor(v))).date(); t = datetime.date.today()
    sunday = t - datetime.timedelta(days=(t.weekday() + 1) % 7)
    ym = lambda x: x.year * 12 + x.month
    return {"today": d == t, "yesterday": d == t - datetime.timedelta(1), "tomorrow": d == t + datetime.timedelta(1),
            "last7Days": t - datetime.timedelta(6) <= d <= t, "thisWeek": sunday <= d < sunday + datetime.timedelta(7),
            "lastWeek": sunday - datetime.timedelta(7) <= d < sunday, "nextWeek": sunday + datetime.timedelta(7) <= d < sunday + datetime.timedelta(14),
            "thisMonth": ym(d) == ym(t), "lastMonth": ym(d) == ym(t) - 1, "nextMonth": ym(d) == ym(t) + 1}.get(p, False)


def lerp(a, b, t):
    pa = [int(a[i:i + 2], 16) for i in (0, 2, 4)]; pb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "#" + "".join(f"{round(x + (y - x) * t):02x}" for x, y in zip(pa, pb))


def rgb(c):
    if c is None or c.type != "rgb" or not isinstance(c.rgb, str):
        return None
    return c.rgb[-6:].lower()


def eval_rule(ctx, ws, cf_ranges, r):
    cells = []
    seen = set()
    for rng in cf_ranges:
        c0, r0, c1, r1 = range_boundaries(rng)
        for row in range(r0, r1 + 1):
            for col in range(c0, c1 + 1):
                ref = f"{get_column_letter(col)}{row}"
                if ref in seen:
                    continue
                seen.add(ref); cells.append((ref, to_val(ws[ref])))
    origin = cf_ranges[0].split(":")[0]
    nums = sorted(v for _, v in cells if isinstance(v, (int, float)) and not isinstance(v, bool) and not isinstance(v, Err))
    out = {}
    t = r.type
    if t == "cellIs":
        opmap = {"greaterThan": ">", "lessThan": "<", "greaterThanOrEqual": ">=", "lessThanOrEqual": "<=", "equal": "=", "notEqual": "<>"}
        fs = list(r.formula or [])

        def operand(i, ref):
            k = const(fs[i])
            return k if k is not None else ctx.evaluate(ws.title, fs[i], origin, ref)
        for ref, v in cells:
            if r.operator in opmap:
                ok = cmp(opmap[r.operator], v, operand(0, ref))
            elif r.operator in ("between", "notBetween"):
                a, b = operand(0, ref), operand(1, ref)
                if isinstance(a, (int, float)) and isinstance(b, (int, float)) and a > b:
                    a, b = b, a
                inside = cmp(">=", v, a) and cmp("<=", v, b)
                ok = inside if r.operator == "between" else (not isinstance(v, Err) and not inside)
            else:
                raise ValueError("operator " + str(r.operator))
            if ok:
                out[ref] = True
    elif t in ("containsText", "notContainsText", "beginsWith", "endsWith"):
        if r.text is None:
            raise ValueError("no text")
        n = r.text.lower()
        for ref, v in cells:
            s = excel_text(v).lower(); e = isinstance(v, Err)
            ok = {"containsText": not e and n in s, "notContainsText": e or n not in s, "beginsWith": not e and s.startswith(n), "endsWith": not e and s.endswith(n)}[t]
            if ok:
                out[ref] = True
    elif t in ("containsBlanks", "notContainsBlanks", "containsErrors", "notContainsErrors"):
        for ref, v in cells:
            e = isinstance(v, Err); blank = not e and excel_text(v).strip(" ") == ""
            if {"containsBlanks": blank, "notContainsBlanks": not e and not blank, "containsErrors": e, "notContainsErrors": not e}[t]:
                out[ref] = True
    elif t in ("duplicateValues", "uniqueValues"):
        def key(v):
            if v is None or v == "" or isinstance(v, Err):
                return None
            return ("s", v.lower()) if isinstance(v, str) else ("v", float(v) if not isinstance(v, bool) else ("b", v))
        counts = {}
        for _, v in cells:
            k = key(v)
            if k is not None:
                counts[k] = counts.get(k, 0) + 1
        for ref, v in cells:
            k = key(v)
            if k is not None and ((counts[k] > 1) == (t == "duplicateValues")):
                out[ref] = True
    elif t == "top10":
        rk = int(r.rank or 10)
        k = max(1, (len(nums) * min(rk, 100)) // 100) if r.percent else rk
        order = nums if r.bottom else nums[::-1]
        if order:
            cut = order[min(k, len(order)) - 1]
            for ref, v in cells:
                if isinstance(v, (int, float)) and not isinstance(v, bool) and (v <= cut if r.bottom else v >= cut):
                    out[ref] = True
    elif t == "aboveAverage":
        if nums:
            avg = sum(nums) / len(nums)
            sd = (statistics.stdev(nums) if len(nums) > 1 else 0) * (r.stdDev or 0)
            above = r.aboveAverage is not False
            comp = avg + sd if above else avg - sd
            for ref, v in cells:
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    ok = (v >= comp if r.equalAverage else v > comp) if above else (v <= comp if r.equalAverage else v < comp)
                    if ok:
                        out[ref] = True
    elif t == "timePeriod":
        for ref, v in cells:
            if period_match(r.timePeriod, v):
                out[ref] = True
    elif t == "expression":
        f = (r.formula or [None])[0]
        if not f:
            raise ValueError("no formula")
        for ref, v in cells:
            if truthy(ctx.evaluate(ws.title, f, origin, ref)):
                out[ref] = True
    elif t in ("colorScale", "dataBar", "iconSet"):
        if not nums:
            return out
        lo, hi = nums[0], nums[-1]

        def cfvo(o):
            if o.type == "min":
                return lo
            if o.type == "max":
                return hi
            val = o.val
            x = float(val) if isinstance(val, (int, float)) else const(str(val)) if val is not None else None
            if x is None or isinstance(x, str):
                x = ctx.evaluate(ws.title, str(val).lstrip("="), origin, origin)
            x = float(x)
            if o.type == "percent":
                return lo + (hi - lo) * x / 100
            if o.type == "percentile":
                return percentile_inc(nums, x)
            return x
        if t == "colorScale":
            cs = r.colorScale; cols = [rgb(c) for c in cs.color]
            if None in cols:
                raise ValueError("non-rgb colour")
            at = [cfvo(o) for o in cs.cfvo]
            for i in range(1, len(at)):
                at[i] = max(at[i], at[i - 1])
            for ref, v in cells:
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    continue
                if v <= at[0]:
                    c = "#" + cols[0]
                elif v >= at[-1]:
                    c = "#" + cols[-1]
                else:
                    c = None
                    for i in range(1, len(at)):
                        if v <= at[i]:
                            c = lerp(cols[i - 1], cols[i], (v - at[i - 1]) / (at[i] - at[i - 1])) if at[i] > at[i - 1] else "#" + cols[i]
                            break
                out[ref] = c
        elif t == "dataBar":
            db = r.dataBar
            a, b = cfvo(db.cfvo[0]), cfvo(db.cfvo[1])
            b = max(a, b)
            mn = (db.minLength if db.minLength is not None else 10) / 100; mx = (db.maxLength if db.maxLength is not None else 90) / 100
            for ref, v in cells:
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    fr = min(1, max(0, (v - a) / (b - a))) if b > a else (1 if v >= b else 0)
                    out[ref] = round(mn + fr * (mx - mn), 6)
        else:
            ic = r.iconSet; name = ic.iconSet or "3TrafficLights1"; n = ICON_N.get(name)
            if n is None:
                raise ValueError("icon set " + str(name))
            vos = list(ic.cfvo or [])
            if len(vos) == n:
                th = [(cfvo(o), o.gte is False) for o in vos]
            else:
                th = [(lo + (hi - lo) * p / 100, False) for p in DEFAULT_PCT[n]]
            for ref, v in cells:
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    continue
                idx = 0
                for i in range(n - 1, 0, -1):
                    tv, gt = th[i]
                    if (v > tv) if gt else (v >= tv):
                        idx = i; break
                if ic.reverse:
                    idx = n - 1 - idx
                out[ref] = idx
    else:
        raise ValueError("type " + str(t))
    return out


for line in sys.stdin:
    p = line.rstrip("\n")
    if not p:
        continue
    try:
        order = doc_order(p)
        wb = openpyxl.load_workbook(p, data_only=True)
        ctx = Ctx(wb); rules = []
        for ws in wb.worksheets:
            for cf in ws.conditional_formatting:
                for r in cf.rules:
                    idx, sq = order.get(ws.title, {}).get(int(r.priority or 0), (None, ""))
                    # openpyxl keeps sqref ranges in a set (file order lost); relative formulas anchor on the FIRST range
                    rngs = sq.replace("$", "").split() or [str(x) for x in cf.sqref.ranges]
                    rec = {"sheet": ws.title, "idx": idx, "priority": r.priority, "type": r.type, "sqref": " ".join(rngs)}
                    try:
                        rec["cells"] = eval_rule(ctx, ws, rngs, r)
                    except Exception as e:  # unsupported by the oracle → excluded from the comparison, reported
                        rec["skip"] = f"{type(e).__name__}: {e}"[:160]
                    rules.append(rec)
        print(json.dumps({"file": p, "rules": rules}), flush=True)
    except Exception as e:
        print(json.dumps({"file": p, "error": f"{type(e).__name__}: {e}"[:200]}), flush=True)
