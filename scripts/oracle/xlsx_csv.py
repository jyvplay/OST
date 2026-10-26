"""TT74 X1 oracle (test-only). manifest.json lists {file, sheet, csv, range:[c0,r0,c1,r1] (1-based) | null, expect:[[...]]}.
For each entry: (a) Python's csv module parses Fusion's CSV -> must equal `expect` (Fusion's intended fields: quoting,
CRLF, BOM); (b) openpyxl data_only (cached values Excel saved) over the same range vs Fusion's raw values; (c) openpyxl's
sheet bounds vs Fusion's used range."""
import csv, io, json, math, os, sys, datetime
import warnings
import openpyxl
warnings.filterwarnings("ignore")
from openpyxl.utils.datetime import to_excel


EPOCH = [None]


def norm(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return to_excel(v, EPOCH[0]) if EPOCH[0] is not None else to_excel(v)
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6) / 86400
    if isinstance(v, datetime.timedelta):
        return v.total_seconds() / 86400
    return v


def same(ours, theirs):
    t = norm(theirs)
    if isinstance(t, (int, float)) and not isinstance(t, bool):
        try:
            o = float(ours)
        except ValueError:
            return False
        # openpyxl turns date/time-formatted numbers into datetime objects rounded to the millisecond
        tol = 2e-8 if isinstance(theirs, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)) else 1e-9 * max(1.0, abs(t))
        return o == t or (math.isfinite(o) and abs(o - t) <= tol)
    return str(ours) == str(t)


def main(d):
    man = json.load(open(os.path.join(d, "manifest.json")))
    books = {}
    out = {"entries": 0, "quoting_ok": 0, "quoting_bad": [], "cells": 0, "agree": 0, "disagree": [], "range_ok": 0, "range_bad": [], "oracle_errors": []}
    for e in man:
        out["entries"] += 1
        text = open(os.path.join(d, e["csv"]), encoding="utf-8", newline="").read()
        # an empty line is a row with one empty field (RFC 4180 ambiguity: Python's csv yields [], Excel writes it so)
        rows = [r if r else [""] for r in csv.reader(io.StringIO(text.lstrip("\ufeff"), newline=""))]
        if rows == e["expect"]:
            out["quoting_ok"] += 1
        else:
            out["quoting_bad"].append(e["csv"])
        try:
            if e["file"] not in books:
                books[e["file"]] = openpyxl.load_workbook(e["path"], data_only=True)
            ws = books[e["file"]][e["sheet"]]
            EPOCH[0] = books[e["file"]].epoch
        except Exception as ex:
            out["oracle_errors"].append(f'{e["file"]}:{e["sheet"]}: {type(ex).__name__}')
            continue
        theirs_range = None
        if not hasattr(ws, "iter_rows"):  # a chartsheet holds no cells
            if e["range"] is None:
                out["range_ok"] += 1
            else:
                out["range_bad"].append(f'{e["file"]}:{e["sheet"]} chartsheet ours={e["range"]}')
            continue
        # strict axis: the extent of value-bearing cells; info axis: extent including formatted-only cells (Excel's
        # UsedRange; Fusion keeps a formatted empty cell only when its formatting is one the cell model holds)
        def bounds(cells):
            return [min(c for _, c in cells), min(r for r, _ in cells), max(c for _, c in cells), max(r for r, _ in cells)] if cells else None
        theirs_v = bounds([(c.row, c.column) for row in ws.iter_rows() for c in row if c.value is not None and c.value != ""])
        theirs_s = bounds([(c.row, c.column) for row in ws.iter_rows() for c in row if c.value is not None or c.has_style])
        ours_v = None
        if e["range"]:
            c0, r0 = e["range"][0], e["range"][1]
            ours_v = bounds([(r0 + i, c0 + j) for i, row in enumerate(e["expect"]) for j, f in enumerate(row) if f != ""])
        if theirs_v == ours_v:
            out["range_ok"] += 1
        elif len(out["range_bad"]) < 40:
            out["range_bad"].append(f'{e["file"]}:{e["sheet"]} values ours={ours_v} theirs={theirs_v}')
        out["styled_range_same"] = out.get("styled_range_same", 0) + (1 if theirs_s == e["range"] else 0)
        if not e["range"]:
            continue
        c0, r0, c1, r1 = e["range"]
        for i, r in enumerate(range(r0, r1 + 1)):
            for j, c in enumerate(range(c0, c1 + 1)):
                ours = e["expect"][i][j] if i < len(e["expect"]) and j < len(e["expect"][i]) else ""
                theirs = ws.cell(row=r, column=c).value
                if ours == "" and theirs is None:
                    continue
                out["cells"] += 1
                if same(ours, theirs):
                    out["agree"] += 1
                elif len(out["disagree"]) < 400:
                    out["disagree"].append([e["file"], e["sheet"], f"{openpyxl.utils.get_column_letter(c)}{r}", ours, str(norm(theirs))[:60]])
    print(json.dumps(out, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])
