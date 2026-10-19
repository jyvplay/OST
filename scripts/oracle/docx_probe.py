"""TT46 independent oracle (python-docx 1.2, MIT): describes a .docx as JSON so tests can verify what our writer
produced without using our own reader. argv: paths. stdout: one JSON object per path."""
import sys, json
import docx
from docx.oxml.ns import qn

def probe(path):
    d = docx.Document(path)
    paras = []
    for p in d.paragraphs:
        ppr = p._p.pPr
        num = ppr.numPr if ppr is not None else None
        nid = num.numId.val if num is not None and num.numId is not None else None
        ilvl = num.ilvl.val if num is not None and num.ilvl is not None else None
        paras.append({
            "style": p.style.name if p.style is not None else None, "text": p.text,
            "align": str(p.alignment) if p.alignment is not None else None,
            "numId": nid, "ilvl": ilvl,
            "runs": [{"text": r.text, "b": r.bold, "i": r.italic, "u": bool(r.underline), "size": r.font.size.pt if r.font.size else None, "font": r.font.name,
                      "color": str(r.font.color.rgb) if r.font.color is not None and r.font.color.type is not None else None} for r in p.runs],
            "links": [{"text": h.text, "url": h.address, "fragment": h.fragment} for h in p.hyperlinks],
            "pageBreak": any(br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br"))),
            "ins": ["".join(t.text or "" for t in i.iter(qn("w:t"))) for i in p._p.iter(qn("w:ins"))],
            "del": ["".join(t.text or "" for t in i.iter(qn("w:delText"))) for i in p._p.iter(qn("w:del"))],
        })
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    shapes = [{"w": s.width, "h": s.height} for s in d.inline_shapes]
    comments = [{"author": c.author, "text": c.text, "initials": c.initials} for c in d.comments] if hasattr(d, "comments") else []
    secs = [{"w": s.page_width, "h": s.page_height, "orient": str(s.orientation), "top": s.top_margin, "left": s.left_margin} for s in d.sections]
    numbering = None
    try:
        np = d.part.numbering_part
        numbering = {"nums": len(np.element.findall(qn("w:num"))), "abstract": len(np.element.findall(qn("w:abstractNum"))),
                     "fmts": [l.find(qn("w:numFmt")).get(qn("w:val")) for a in np.element.findall(qn("w:abstractNum")) for l in a.findall(qn("w:lvl"))[:1] if l.find(qn("w:numFmt")) is not None]}
    except Exception:
        pass
    return {"file": path, "paras": paras, "tables": tables, "shapes": shapes, "comments": comments, "sections": secs, "numbering": numbering}

for a in sys.argv[1:]:
    try:
        print(json.dumps(probe(a)))
    except Exception as e:
        print(json.dumps({"file": a, "error": f"{type(e).__name__}: {e}"}))
