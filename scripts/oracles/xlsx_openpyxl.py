"""Independent OOXML oracle using openpyxl 3.1 (PyPI).

  read <file.xlsx>   -> JSON dump of what openpyxl sees in OUR writer's output
  write <file.xlsx>  -> author a styled workbook with openpyxl for OUR reader
                        (includes a shared formula block and an array formula,
                        injected at the XML level because openpyxl does not
                        author shared formulas itself)
"""
import json, sys, zipfile, shutil, os, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def rgb(c):
    if c is None or c.type != "rgb" or c.rgb in (None, "00000000"):
        return None
    return "#" + str(c.rgb)[-6:].lower()


def read(path):
    wb = openpyxl.load_workbook(path)  # formulas as text
    wbv = openpyxl.load_workbook(path, data_only=True)  # cached values
    out = {"sheets": []}
    for ws in wb.worksheets:
        wv = wbv[ws.title]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                f = c.font
                v = c.value
                if hasattr(v, "text"):  # ArrayFormula
                    v = {"array": v.text, "ref": v.ref}
                rec = {
                    "value": v if isinstance(v, (int, float, str, bool, dict)) or v is None else str(v),
                    "cached": wv[c.coordinate].value if not hasattr(wv[c.coordinate].value, "isoformat") else str(wv[c.coordinate].value),
                    "bold": bool(f.b), "italic": bool(f.i), "underline": f.u,
                    "color": rgb(f.color) if f.color is not None else None,
                    "fill": rgb(c.fill.fgColor) if c.fill.fill_type == "solid" else None,
                    "numFmt": c.number_format,
                    "align": c.alignment.horizontal, "wrap": bool(c.alignment.wrap_text),
                }
                if c.value is None and not (rec["bold"] or rec["fill"] or rec["numFmt"] != "General"):
                    continue
                cells[c.coordinate] = rec
        widths = {k: d.width for k, d in ws.column_dimensions.items() if d.customWidth}
        out["sheets"].append({"title": ws.title, "cells": cells, "widths": widths, "freeze": ws.freeze_panes})
    print(json.dumps(out, default=str))


def layout_read(path):
    """Merges, borders, vertical alignment, data validations and panes as openpyxl sees them."""
    wb = openpyxl.load_workbook(path)
    out = {"sheets": []}
    for ws in wb.worksheets:
        borders = {}
        valign = {}
        for row in ws.iter_rows():
            for c in row:
                b = c.border
                sides = {}
                for k in ("left", "right", "top", "bottom"):
                    sd = getattr(b, k)
                    if sd is not None and sd.style:
                        sides[k] = {"style": sd.style, "color": rgb(sd.color) if sd.color is not None else None}
                if sides:
                    borders[c.coordinate] = sides
                if c.alignment.vertical:
                    valign[c.coordinate] = c.alignment.vertical
        dvs = []
        for dv in ws.data_validations.dataValidation:
            dvs.append({"type": dv.type, "operator": dv.operator, "sqref": str(dv.sqref), "formula1": dv.formula1, "formula2": dv.formula2,
                        "allowBlank": bool(dv.allow_blank), "showDropDown": bool(dv.showDropDown), "showErrorMessage": bool(dv.showErrorMessage),
                        "errorStyle": dv.errorStyle, "errorTitle": dv.errorTitle, "error": dv.error, "promptTitle": dv.promptTitle, "prompt": dv.prompt})
        out["sheets"].append({"title": ws.title, "merges": sorted(str(r) for r in ws.merged_cells.ranges), "borders": borders, "valign": valign,
                              "validations": dvs, "freeze": ws.freeze_panes})
    print(json.dumps(out, default=str))


def layout_write(path):
    from openpyxl.styles import Border, Side
    from openpyxl.styles.colors import Color
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Layout"
    ws["A1"] = "Quarterly report"; ws["C1"] = "lost"  # merged A1:D1 — openpyxl turns C1 into a MergedCell (value dropped), like Excel
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A3:A5")
    thin = Side(style="thin", color="FF000000")
    ws["B3"].border = Border(left=thin, right=Side(style="medium", color="FFFF0000"), top=Side(style="dashed"), bottom=Side(style="double", color="FF00B050"))
    ws["C3"].border = Border(bottom=Side(style="thick", color=Color(theme=4)))  # theme colour: style kept, colour unset
    ws["D3"].border = Border(left=Side(style="mediumDashed"))  # style outside our set: dropped (declared)
    ws["B4"].alignment = Alignment(vertical="top")
    ws["F1"] = "Red"; ws["F2"] = "Green"; ws["F3"] = "Blue"
    dv1 = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=True); dv1.add("B6:B20")
    dv2 = DataValidation(type="list", formula1="$F$1:$F$3", allow_blank=False, showDropDown=True); dv2.add("C6:C20")  # showDropDown=True HIDES the arrow
    dv3 = DataValidation(type="whole", operator="between", formula1="1", formula2="10", errorStyle="warning", errorTitle="Range", error="1 to 10 please"); dv3.add("D6:D20"); dv3.add("E6")
    dv4 = DataValidation(type="decimal", operator="greaterThan", formula1="0", showErrorMessage=True, promptTitle="Price", prompt="positive"); dv4.add("G6:G20")
    dv5 = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="5", showErrorMessage=True); dv5.add("H6:H20")
    dv6 = DataValidation(type="custom", formula1="MOD(I6,2)=0", showErrorMessage=True); dv6.add("I6:I20")
    for dv in (dv1, dv2, dv3, dv4, dv5, dv6):
        if dv not in (dv3,):
            dv.showErrorMessage = True
        ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    wb.save(path)
    print(json.dumps({"ok": True}))


def write(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Item"; ws["A1"].font = Font(bold=True, italic=True, color="FF1F4E79")
    ws["B1"] = "Price"; ws["B1"].fill = PatternFill("solid", fgColor="FFFFF2CC")
    ws["B1"].alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (n, p) in enumerate([("a", 1.5), ("b", 2.25), ("c", 4)], start=2):
        ws[f"A{i}"] = n; ws[f"B{i}"] = p; ws[f"B{i}"].number_format = "$#,##0.00"
        ws[f"C{i}"] = f"=B{i}*2"  # replaced by a shared formula block below
    ws["D2"] = "=TEXTJOIN(\",\",TRUE,A2:A4)"  # openpyxl writes it verbatim; we prefix in XML below
    ws["E1"].font = Font(underline="single")  # styled blank
    ws["F2"] = 45361; ws["F2"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 20
    ws.freeze_panes = "A2"
    tmp = path + ".tmp"
    wb.save(tmp)
    # inject Excel-style shared + array formulas and the _xlfn prefix at the XML level
    zin = zipfile.ZipFile(tmp)
    zout = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "xl/worksheets/sheet1.xml":
            x = data.decode()
            x = x.replace('<f>B2*2</f>', '<f t="shared" ref="C2:C4" si="0">B2*2</f>')
            x = x.replace('<f>B3*2</f>', '<f t="shared" si="0"/>').replace('<f>B4*2</f>', '<f t="shared" si="0"/>')
            x = x.replace('TEXTJOIN(', '_xlfn.TEXTJOIN(')
            # array formula at G2:G4 with stale cached children that must not block the re-spill
            x = x.replace('</row><row r="3"', '<c r="G2"><f t="array" ref="G2:G4">B2:B4*10</f><v>15</v></c></row><row r="3"', 1)
            x = x.replace('</row><row r="4"', '<c r="G3"><v>999</v></c></row><row r="4"', 1)
            x = x.replace('</row></sheetData>', '<c r="G4"><v>999</v></c></row></sheetData>', 1)
            data = x.encode()
        zout.writestr(item, data)
    zout.close(); zin.close(); os.remove(tmp)
    print(json.dumps({"ok": True}))


if __name__ == "__main__":
    {"read": read, "write": write, "layout-read": layout_read, "layout-write": layout_write}[sys.argv[1]](sys.argv[2])
