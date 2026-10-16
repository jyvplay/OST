"""Independent oracle for the TT18 metadata cleaner (src/lib/privacy/metaclean.ts).

  make <dir>             author metadata-laden fixtures with Pillow / python-docx / openpyxl / python-pptx,
                         plus hand-injected segments (XMP, Photoshop IRB, GIF XMP app-ext, docx rsids,
                         docProps/custom.xml, xlsx absPath) that those libraries do not write themselves
  verify <dir> <clean>   for each cleaned file: decode original and cleaned with the SAME independent
                         library and require identical pixels/text, then require the metadata to be gone
"""
import io, json, os, re, struct, sys, zipfile

from PIL import Image, PngImagePlugin

PERSON = "Jane Doe"


def gradient(w=64, h=40, mode="RGB"):
    im = Image.new(mode, (w, h))
    px = im.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = ((x * 4) % 256, (y * 6) % 256, ((x + y) * 3) % 256) if mode == "RGB" else ((x * 4) % 256, (y * 6) % 256, 90, 255)
    return im


def jpeg_insert(data, seg):
    return data[:2] + seg + data[2:]


def app(marker, payload):
    return bytes([0xFF, marker]) + struct.pack(">H", len(payload) + 2) + payload


def make(d):
    os.makedirs(d, exist_ok=True)
    # JPEG: EXIF (camera, author, date, software, GPS, orientation 6) + COM + XMP + Photoshop IRB
    ex = Image.Exif()
    ex[0x010F] = "Canon"; ex[0x0110] = "EOS R5"; ex[0x0132] = "2026:10:16 09:30:00"; ex[0x013B] = PERSON
    ex[0x0131] = "PhotoEditor 9"; ex[0x0112] = 6
    gps = ex.get_ifd(0x8825)
    gps[1] = "N"; gps[2] = (41.0, 30.0, 0.0); gps[3] = "W"; gps[4] = (81.0, 31.0, 0.0)
    b = io.BytesIO(); gradient().save(b, "JPEG", quality=92, exif=ex, comment=b"shot by Jane at home")
    data = b.getvalue()
    xmp = b"http://ns.adobe.com/xap/1.0/\0<x:xmpmeta><dc:creator>Jane Doe</dc:creator></x:xmpmeta>"
    data = jpeg_insert(data, app(0xE1, xmp))
    data = jpeg_insert(data, app(0xED, b"Photoshop 3.0\08BIM\x04\x04\0\0\0\0\0\x0aJane Doe!!"))
    open(f"{d}/photo.jpg", "wb").write(data)
    # PNG: tEXt / zTXt / iTXt / eXIf + a tIME chunk injected after IHDR
    info = PngImagePlugin.PngInfo()
    info.add_text("Author", PERSON); info.add_text("Software", "Painter 3"); info.add_text("Comment", "draft " * 30, zip=True)
    info.add_itxt("Description", "secret plan", lang="en", tkey="Description")
    pe = Image.Exif(); pe[0x013B] = PERSON
    b = io.BytesIO(); gradient(mode="RGBA").save(b, "PNG", pnginfo=info, exif=pe)
    data = b.getvalue()
    t = b"tIME" + bytes([0x07, 0xEA, 10, 16, 9, 30, 0])
    import zlib
    chunk = struct.pack(">I", 7) + t + struct.pack(">I", zlib.crc32(t) & 0xFFFFFFFF)
    data = data[:33] + chunk + data[33:]
    open(f"{d}/pic.png", "wb").write(data)
    # WebP lossless with EXIF + XMP
    we = Image.Exif(); we[0x013B] = PERSON; we[0x010F] = "Pixel"
    b = io.BytesIO(); gradient(mode="RGBA").save(b, "WEBP", lossless=True, exif=we, xmp=b"<x:xmpmeta>Jane Doe</x:xmpmeta>")
    open(f"{d}/img.webp", "wb").write(b.getvalue())
    # GIF: 2 frames, loop, comment + an XMP application extension injected after the colour table
    f1 = gradient().convert("P"); f2 = gradient().transpose(Image.FLIP_LEFT_RIGHT).convert("P")
    b = io.BytesIO(); f1.save(b, "GIF", save_all=True, append_images=[f2], loop=0, duration=100, comment=b"made by Jane Doe")
    data = b.getvalue()
    flags = data[10]
    p = 13 + (3 * (1 << ((flags & 7) + 1)) if flags & 0x80 else 0)
    payload = b"<x:xmpmeta>Jane Doe</x:xmpmeta>"
    xmp_ext = b"\x21\xff\x0bXMP DataXMP" + bytes([len(payload)]) + payload + b"\x00"
    data = data[:p] + xmp_ext + data[p:]
    open(f"{d}/anim.gif", "wb").write(data)
    # DOCX: core props, a comment, then rsids + custom.xml injected at XML level
    import docx
    doc = docx.Document()
    doc.core_properties.author = PERSON; doc.core_properties.last_modified_by = PERSON
    doc.core_properties.title = "Quarterly plan"; doc.core_properties.comments = "internal only"
    par = doc.add_paragraph("The visible text must survive cleaning.")
    doc.add_paragraph("Second paragraph.")
    try:
        doc.add_comment(par.runs, text="check this", author=PERSON, initials="JD")
    except Exception:
        pass
    b = io.BytesIO(); doc.save(b)
    open(f"{d}/doc.docx", "wb").write(inject_docx(b.getvalue()))
    # XLSX: creator + legacy comment + absPath/revisionPtr injected into workbook.xml
    import openpyxl
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws["A1"] = "Total"; ws["B1"] = 42; ws["A2"] = "=B1*2"
    ws["B1"].comment = Comment("verify this", PERSON)
    wb.properties.creator = PERSON; wb.properties.lastModifiedBy = PERSON
    b = io.BytesIO(); wb.save(b)
    open(f"{d}/book.xlsx", "wb").write(inject_xlsx(b.getvalue()))
    # PPTX
    import pptx
    pr = pptx.Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Roadmap"; s.placeholders[1].text = "Ship privacy"
    pr.core_properties.author = PERSON; pr.core_properties.last_modified_by = PERSON
    b = io.BytesIO(); pr.save(b)
    open(f"{d}/deck.pptx", "wb").write(b.getvalue())
    # SVG from an editor
    open(f"{d}/draw.svg", "w").write(
        '<?xml version="1.0"?>\n<!-- Created by Jane Doe with Inkscape -->\n'
        '<svg xmlns="http://www.w3.org/2000/svg" xmlns:inkscape="http://www.inkscape.org/namespaces/inkscape" '
        'xmlns:sodipodi="http://sodipodi.sourceforge.net/DTD/sodipodi-0.dtd" xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" width="100" height="60" sodipodi:docname="C:\\Users\\jane\\plan.svg" inkscape:version="1.3">'
        '<sodipodi:namedview id="nv" inkscape:zoom="2" pagecolor="#fff"/>'
        '<metadata><rdf:RDF><dc:creator>Jane Doe</dc:creator></rdf:RDF></metadata>'
        '<title>Plan</title><rect x="5" y="5" width="40" height="20" fill="#36c" inkscape:label="box"/>'
        '<circle cx="70" cy="30" r="15" fill="#c63"/><path d="M5 50 L95 50" stroke="#000"/></svg>')
    return {"files": sorted(os.listdir(d))}


def inject_docx(data):
    zin = zipfile.ZipFile(io.BytesIO(data)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        x = zin.read(it.filename)
        if it.filename == "word/document.xml":
            x = x.replace(b"<w:p>", b'<w:p w:rsidR="00AB12CD" w:rsidRDefault="00EF3456">').replace(b"<w:r>", b'<w:r w:rsidRPr="00AB12CD">')
        if it.filename == "word/settings.xml":
            x = x.replace(b"</w:settings>", b'<w:rsids><w:rsidRoot w:val="00AB12CD"/><w:rsid w:val="00EF3456"/></w:rsids></w:settings>')
        if it.filename == "[Content_Types].xml":
            x = x.replace(b"</Types>", b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
        if it.filename == "_rels/.rels":
            x = x.replace(b"</Relationships>", b'<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        zout.writestr(it, x)
    zout.writestr("docProps/custom.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Client"><vt:lpwstr>ACME secret client</vt:lpwstr></property></Properties>')
    zout.close(); return out.getvalue()


def inject_xlsx(data):
    zin = zipfile.ZipFile(io.BytesIO(data)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        x = zin.read(it.filename)
        if it.filename == "xl/workbook.xml":
            x = re.sub(rb"(<workbook\b[^>]*)>", rb'\1 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac" xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" mc:Ignorable="x15ac xr">'
                       rb'<mc:AlternateContent><mc:Choice Requires="x15"><x15ac:absPath url="C:\\Users\\jane\\Documents\\Finance\\"/></mc:Choice></mc:AlternateContent>'
                       rb'<xr:revisionPtr revIDLastSave="0" documentId="13_ncr:1_{5B3F7F0E-1111-2222-3333-444455556666}"/>', x, count=1)
        zout.writestr(it, x)
    zout.close(); return out.getvalue()


def png_chunks(data):
    p, out = 8, []
    while p < len(data):
        n = struct.unpack(">I", data[p:p + 4])[0]; out.append(data[p + 4:p + 8].decode()); p += 12 + n
    return out


def jpeg_markers(data):
    p, out = 2, []
    while p < len(data):
        m = data[p + 1]
        if m == 0xDA:
            out.append("SOS"); break
        n = struct.unpack(">H", data[p + 2:p + 4])[0]; out.append(hex(m)); p += 2 + n
    return out


def frames(im):
    out = []
    try:
        while True:
            out.append(im.convert("RGBA").tobytes()); im.seek(im.tell() + 1)
    except EOFError:
        pass
    return out


def verify(d, c):
    res = {}
    for name in sorted(os.listdir(c)):
        o, k = open(f"{d}/{name}", "rb").read(), open(f"{c}/{name}", "rb").read()
        probs, info = [], {}
        try:
            if name.endswith((".jpg", ".png", ".webp", ".gif")):
                a, b = Image.open(io.BytesIO(o)), Image.open(io.BytesIO(k))
                if frames(a) != frames(Image.open(io.BytesIO(k))):
                    probs.append("decoded pixels differ")
                b = Image.open(io.BytesIO(k))
                exif = dict(b.getexif())
                info["exif"] = {str(t): str(v) for t, v in exif.items()}
                info["info_keys"] = sorted(str(x) for x in b.info.keys())
                txt = json.dumps(info) + k.decode("latin1")
                if PERSON in txt or "Canon" in txt.split("exif")[0] or "secret plan" in txt:
                    probs.append("identifying text still present in the bytes")
                if name.endswith(".jpg"):
                    if exif != {274: 6}:
                        probs.append(f"EXIF should be exactly Orientation=6, got {exif}")
                    mk = jpeg_markers(k); info["markers"] = mk
                    if any(m in mk for m in ("0xfe", "0xed")):
                        probs.append("COM or APP13 left")
                    if b.getexif().get_ifd(0x8825):
                        probs.append("GPS left")
                if name.endswith(".png"):
                    ch = png_chunks(k); info["chunks"] = ch
                    if set(ch) & {"tEXt", "zTXt", "iTXt", "eXIf", "tIME"}:
                        probs.append(f"text/time chunks left: {ch}")
                if name.endswith(".webp") and ("exif" in b.info or "xmp" in b.info):
                    probs.append("webp exif/xmp left")
                if name.endswith(".gif"):
                    if "comment" in b.info:
                        probs.append("gif comment left")
                    if b.info.get("loop") != 0:
                        probs.append(f"loop lost: {b.info.get('loop')}")
            elif name.endswith((".docx", ".xlsx", ".pptx")):
                z = zipfile.ZipFile(io.BytesIO(k))
                info["parts"] = sorted(z.namelist())
                if any(i.date_time != (1980, 1, 1, 0, 0, 0) for i in z.infolist()):
                    probs.append("zip timestamps not normalised")
                if any(i.extra for i in z.infolist()):
                    probs.append("zip extra fields left")
                allxml = "".join(z.read(n).decode("utf8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))
                for bad in (PERSON, "w:rsid", "absPath", "ACME", "revisionPtr", "Quarterly plan", "internal only"):
                    if bad in allxml:
                        probs.append(f"{bad!r} still present")
                if name.endswith(".docx"):
                    import docx
                    A, B = docx.Document(io.BytesIO(o)), docx.Document(io.BytesIO(k))
                    if [p.text for p in A.paragraphs] != [p.text for p in B.paragraphs]:
                        probs.append("paragraph text changed")
                    cp = B.core_properties
                    info["core"] = {"author": cp.author, "last_modified_by": cp.last_modified_by, "title": cp.title}
                    if cp.author or cp.last_modified_by or cp.title:
                        probs.append(f"core properties left {info['core']}")
                    if "docProps/custom.xml" in z.namelist() or "custom.xml" in z.read("[Content_Types].xml").decode() + z.read("_rels/.rels").decode():
                        probs.append("custom.xml or its references left")
                if name.endswith(".xlsx"):
                    import openpyxl
                    A, B = openpyxl.load_workbook(io.BytesIO(o)), openpyxl.load_workbook(io.BytesIO(k))
                    va = [[c.value for c in r] for r in A.active.iter_rows()]; vb = [[c.value for c in r] for r in B.active.iter_rows()]
                    if va != vb:
                        probs.append("cell values changed")
                    # openpyxl SUBSTITUTES creator="openpyxl" when core.xml has no dc:creator, so the
                    # library property cannot prove absence; check the raw part instead.
                    core = z.read("docProps/core.xml").decode()
                    info["creator"] = B.properties.creator
                    if re.search(r"<dc:creator|<cp:lastModifiedBy", core):
                        probs.append("workbook creator left in docProps/core.xml")
                    cm = B.active["B1"].comment
                    info["comment_author"] = cm.author if cm else None
                    if cm is None or cm.text.strip() == "":
                        probs.append("comment text lost (only the author should go)")
                if name.endswith(".pptx"):
                    import pptx
                    A, B = pptx.Presentation(io.BytesIO(o)), pptx.Presentation(io.BytesIO(k))
                    ta = [sh.text_frame.text for s in A.slides for sh in s.shapes if sh.has_text_frame]
                    tb = [sh.text_frame.text for s in B.slides for sh in s.shapes if sh.has_text_frame]
                    if ta != tb:
                        probs.append("slide text changed")
                    if B.core_properties.author or B.core_properties.last_modified_by:
                        probs.append("pptx author left")
            elif name.endswith(".svg"):
                import svgelements
                A = list(svgelements.SVG.parse(io.BytesIO(o)).elements()); B = list(svgelements.SVG.parse(io.BytesIO(k)).elements())
                sa = [type(e).__name__ for e in A if isinstance(e, svgelements.Shape)]; sb = [type(e).__name__ for e in B if isinstance(e, svgelements.Shape)]
                info["shapes"] = sb
                if sa != sb:
                    probs.append(f"shapes changed {sa} vs {sb}")
                t = k.decode()
                for bad in ("inkscape:", "sodipodi:", "<metadata", "<!--", PERSON, "C:\\Users"):
                    if bad in t:
                        probs.append(f"{bad!r} left")
                if "<title>Plan</title>" not in t:
                    probs.append("title (content) was removed")
        except Exception as e:  # a crash is a failure, reported not raised
            probs.append(f"oracle could not open cleaned file: {type(e).__name__}: {e}")
        res[name] = {"problems": probs, "info": info}
    return res


if __name__ == "__main__":
    if sys.argv[1] == "make":
        print(json.dumps(make(sys.argv[2])))
    else:
        print(json.dumps(verify(sys.argv[2], sys.argv[3]), default=str))
