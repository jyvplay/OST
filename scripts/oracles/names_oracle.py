"""Independent oracle for TT19 defined names.

  make <path>      openpyxl-authored workbook with workbook-scope, sheet-scope, formula and constant names,
                   plus a hidden name and _xlnm.Print_Area (must be skipped), and formulas that use the names.
                   Prints the expected values computed here in Python (not by our engine).
  inspect <path>   reads OUR written file with openpyxl: names, scopes, destinations, formulas.
  calc <path>      evaluates OUR written file with the PyPI `formulas` engine (independent calculator) and
                   prints the computed values of Data!D1:D6.
"""
import json, sys

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

SALES = [120, 340, 75.5, 410, 99]
RATE = 0.07


def make(path):
    wb = openpyxl.Workbook()
    data = wb.active; data.title = "Data"
    rates = wb.create_sheet("Rates")
    data["A1"] = "Sales"
    for i, v in enumerate(SALES):
        data.cell(row=i + 2, column=2, value=v)
    rates["B1"] = RATE
    wb.defined_names["Sales"] = DefinedName("Sales", attr_text="Data!$B$2:$B$6")
    wb.defined_names["Tax_Rate"] = DefinedName("Tax_Rate", attr_text="Rates!$B$1")
    wb.defined_names["Grand.Total"] = DefinedName("Grand.Total", attr_text="SUM(Data!$B$2:$B$6)")
    wb.defined_names["Bonus"] = DefinedName("Bonus", attr_text="25", comment="flat bonus")
    wb.defined_names["Secret"] = DefinedName("Secret", attr_text="Data!$A$1", hidden=True)
    data.defined_names["First"] = DefinedName("First", attr_text="Data!$B$2")      # sheet scope (Data)
    rates.defined_names["First"] = DefinedName("First", attr_text="Rates!$B$1")    # same name, other sheet
    data.print_area = "A1:B6"                                                       # _xlnm.Print_Area
    data["D1"] = "=SUM(Sales)"
    data["D2"] = "=SUM(Sales)*Tax_Rate"
    data["D3"] = "=First*2"
    data["D4"] = "=Grand.Total+Bonus"
    data["D5"] = "=COUNT(Sales)"
    data["D6"] = "=INDEX(Sales,3)"
    rates["D1"] = "=First*100"
    wb.save(path)
    s = sum(SALES)
    return {"Data!D1": s, "Data!D2": s * RATE, "Data!D3": SALES[0] * 2, "Data!D4": s + 25, "Data!D5": 5, "Data!D6": SALES[2], "Rates!D1": RATE * 100}


def inspect(path):
    wb = openpyxl.load_workbook(path)
    names = {n: {"ref": d.attr_text, "comment": d.comment, "scope": None} for n, d in wb.defined_names.items()}
    for ws in wb.worksheets:
        for n, d in ws.defined_names.items():
            names[f"{ws.title}!{n}"] = {"ref": d.attr_text, "scope": ws.title}
    dests = {}
    for n, d in wb.defined_names.items():
        try:
            dests[n] = [[t, c] for t, c in d.destinations]
        except Exception as e:  # formula / constant names have no destinations
            dests[n] = f"no destinations: {type(e).__name__}"
    formulas = {f"{ws.title}!{c.coordinate}": c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}
    return {"names": names, "destinations": dests, "formulas": formulas}


def calc(path):
    import formulas
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    out = {}
    for k, v in sol.items():
        ks = str(k).upper()
        for cell in ("D1", "D2", "D3", "D4", "D5", "D6"):
            if ks.endswith(f"[{path.split('/')[-1].upper()}]DATA'!{cell}") or ks.endswith(f"DATA'!{cell}"):
                val = v.value[0][0] if hasattr(v, "value") else v
                out[f"Data!{cell}"] = val if isinstance(val, (int, float, str, bool)) else str(val)
    return out


if __name__ == "__main__":
    op, path = sys.argv[1], sys.argv[2]
    print(json.dumps({"make": make, "inspect": inspect, "calc": calc}[op](path), default=str))
